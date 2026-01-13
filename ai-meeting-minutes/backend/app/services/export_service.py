"""Service for exporting meeting minutes to various formats"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from app.schemas.meeting import MeetingMinutes
from app.core.config import settings


class ExportService:
    """Service for exporting meeting minutes"""
    
    def __init__(self):
        self.export_dir = settings.EXPORT_DIR
        os.makedirs(self.export_dir, exist_ok=True)
    
    def export_to_docx(
        self,
        meeting_id: int,
        minutes: MeetingMinutes,
        title: str = "Meeting Minutes"
    ) -> str:
        """
        Export meeting minutes to DOCX format
        
        Returns:
            Path to exported file
        """
        
        filename = f"meeting_{meeting_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = os.path.join(self.export_dir, filename)
        
        doc = Document()
        
        # Title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Date
        date_para = doc.add_paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()  # Spacing
        
        # Summary
        doc.add_heading("Meeting Summary", level=1)
        doc.add_paragraph(minutes.summary)
        
        # Metadata
        doc.add_heading("Meeting Details", level=1)
        if minutes.duration_minutes:
            doc.add_paragraph(f"Duration: Approximately {minutes.duration_minutes:.0f} minutes")
        if minutes.participants:
            doc.add_paragraph(f"Participants ({len(minutes.participants)}): {', '.join(minutes.participants)}")
        
        # Key Points
        if minutes.key_points:
            doc.add_heading("Key Discussion Points", level=1)
            for point in minutes.key_points:
                doc.add_paragraph(point, style='List Bullet')
        
        # Decisions
        if minutes.decisions:
            doc.add_heading("Decisions Made", level=1)
            for decision in minutes.decisions:
                doc.add_paragraph(decision, style='List Bullet')
        
        # Action Items
        if minutes.action_items:
            doc.add_heading("Action Items", level=1)
            for item in minutes.action_items:
                para = doc.add_paragraph(style='List Bullet')
                para.add_run(item.description)
                if item.owner:
                    para.add_run(f" (Owner: {item.owner})")
                if item.due_date:
                    para.add_run(f" [Due: {item.due_date}]")
        
        # Next Steps
        if minutes.next_steps:
            doc.add_heading("Next Steps", level=1)
            for step in minutes.next_steps:
                doc.add_paragraph(step, style='List Bullet')
        
        doc.save(filepath)
        return filepath
    
    def export_to_pdf(
        self,
        meeting_id: int,
        minutes: MeetingMinutes,
        title: str = "Meeting Minutes"
    ) -> str:
        """
        Export meeting minutes to PDF format
        
        Returns:
            Path to exported file
        """
        
        filename = f"meeting_{meeting_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        filepath = os.path.join(self.export_dir, filename)
        
        doc = SimpleDocTemplate(filepath, pagesize=letter)
        story = []
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=RGBColor(66, 66, 66),
            alignment=TA_CENTER,
            spaceAfter=30
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=RGBColor(88, 88, 88),
            spaceAfter=12,
            spaceBefore=12
        )
        
        # Title
        story.append(Paragraph(title, title_style))
        story.append(Paragraph(
            f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}",
            styles['Normal']
        ))
        story.append(Spacer(1, 0.3*inch))
        
        # Summary
        story.append(Paragraph("Meeting Summary", heading_style))
        story.append(Paragraph(minutes.summary, styles['Normal']))
        story.append(Spacer(1, 0.2*inch))
        
        # Meeting Details
        story.append(Paragraph("Meeting Details", heading_style))
        if minutes.duration_minutes:
            story.append(Paragraph(
                f"<b>Duration:</b> Approximately {minutes.duration_minutes:.0f} minutes",
                styles['Normal']
            ))
        if minutes.participants:
            story.append(Paragraph(
                f"<b>Participants ({len(minutes.participants)}):</b> {', '.join(minutes.participants)}",
                styles['Normal']
            ))
        story.append(Spacer(1, 0.2*inch))
        
        # Key Points
        if minutes.key_points:
            story.append(Paragraph("Key Discussion Points", heading_style))
            for point in minutes.key_points:
                story.append(Paragraph(f"• {point}", styles['Normal']))
            story.append(Spacer(1, 0.2*inch))
        
        # Decisions
        if minutes.decisions:
            story.append(Paragraph("Decisions Made", heading_style))
            for decision in minutes.decisions:
                story.append(Paragraph(f"• {decision}", styles['Normal']))
            story.append(Spacer(1, 0.2*inch))
        
        # Action Items
        if minutes.action_items:
            story.append(Paragraph("Action Items", heading_style))
            for item in minutes.action_items:
                text = f"• {item.description}"
                if item.owner:
                    text += f" <b>(Owner: {item.owner})</b>"
                if item.due_date:
                    text += f" <i>[Due: {item.due_date}]</i>"
                story.append(Paragraph(text, styles['Normal']))
            story.append(Spacer(1, 0.2*inch))
        
        # Next Steps
        if minutes.next_steps:
            story.append(Paragraph("Next Steps", heading_style))
            for step in minutes.next_steps:
                story.append(Paragraph(f"• {step}", styles['Normal']))
        
        doc.build(story)
        return filepath
    
    def export_to_txt(
        self,
        meeting_id: int,
        minutes: MeetingMinutes,
        title: str = "Meeting Minutes"
    ) -> str:
        """Export meeting minutes to plain text format"""
        
        filename = f"meeting_{meeting_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
        filepath = os.path.join(self.export_dir, filename)
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(f"{title}\n")
            f.write("=" * len(title) + "\n\n")
            f.write(f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}\n\n")
            
            f.write("MEETING SUMMARY\n")
            f.write("-" * 50 + "\n")
            f.write(f"{minutes.summary}\n\n")
            
            f.write("MEETING DETAILS\n")
            f.write("-" * 50 + "\n")
            if minutes.duration_minutes:
                f.write(f"Duration: Approximately {minutes.duration_minutes:.0f} minutes\n")
            if minutes.participants:
                f.write(f"Participants ({len(minutes.participants)}): {', '.join(minutes.participants)}\n")
            f.write("\n")
            
            if minutes.key_points:
                f.write("KEY DISCUSSION POINTS\n")
                f.write("-" * 50 + "\n")
                for i, point in enumerate(minutes.key_points, 1):
                    f.write(f"{i}. {point}\n")
                f.write("\n")
            
            if minutes.decisions:
                f.write("DECISIONS MADE\n")
                f.write("-" * 50 + "\n")
                for i, decision in enumerate(minutes.decisions, 1):
                    f.write(f"{i}. {decision}\n")
                f.write("\n")
            
            if minutes.action_items:
                f.write("ACTION ITEMS\n")
                f.write("-" * 50 + "\n")
                for i, item in enumerate(minutes.action_items, 1):
                    f.write(f"{i}. {item.description}")
                    if item.owner:
                        f.write(f" (Owner: {item.owner})")
                    if item.due_date:
                        f.write(f" [Due: {item.due_date}]")
                    f.write("\n")
                f.write("\n")
            
            if minutes.next_steps:
                f.write("NEXT STEPS\n")
                f.write("-" * 50 + "\n")
                for i, step in enumerate(minutes.next_steps, 1):
                    f.write(f"{i}. {step}\n")
        
        return filepath


# Singleton instance
export_service = ExportService()
